"""build_docx —— 把 manuscript.md 排成 Elsevier 模板壳的 .docx。

* A4，2.5 cm 页边距，连续行号，页脚页码
* 公式：LaTeX -> MathML -> OMML（Word 原生公式，可编辑）
* 表格：三线表（仅上下框线 + 表头下框线），8 pt
* 图：600 dpi PNG，图注在下，8 pt
* 样式全部沿用模板的 Els-* 风格；一级/二级标题用模板自带的自动编号
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import latex2mathml.converter as latex2mathml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "manuscript" / "manuscript.md"
OUT = ROOT / "manuscript" / "paper8_manuscript_v1.docx"
TEMPLATE = Path("Elsevier-template.docx")  # local template copy; path sanitised
XSL = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

FIG = {
    1: ("figures/fig1_scheme.png", 15.5),
    2: ("figures/fig1_forward.png", 15.5),
    3: ("figures/fig2_audit.png", 14.0),
    4: ("figures/fig3_extrapolation.png", 15.0),
    5: ("figures/fig4_design.png", 14.0),
    6: ("figures/fig5_robustness.png", 15.5),
}

TBL_WIDTHS = {
    1: [1.22, 2.26, 1.50, 1.60, 1.40, 0.90, 0.95, 1.60, 1.70, 1.50],
    2: [1.20, 1.35, 1.55, 2.60, 0.95, 1.00, 1.40, 1.45, 1.75, 1.80],
    3: [2.40, 2.40, 2.20, 4.20, 4.20],   # (a)
    4: [4.50, 4.50, 4.50],               # (b)
    5: [4.00, 4.00, 4.00, 4.00],         # (c)
    6: [1.30, 1.30, 2.00, 2.20, 2.20, 2.00, 2.40, 2.40],  # table 4
}

_tbl_counter = 0


# ---------------------------------------------------------------------------
# 基础工具
# ---------------------------------------------------------------------------
def omml(tex: str):
    mathml = latex2mathml.convert(tex)
    tree = etree.fromstring(mathml.encode())
    result = etree.XSLT(etree.parse(str(XSL)))(tree)
    xml = etree.tostring(result, encoding="unicode")
    return parse_xml(xml)


TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|\$[^$\n]+?\$)")


def add_inline(p, text: str, size: float | None = None, color=None):
    """写入一段带 **粗体**、*斜体*、$行内公式$ 的文本。"""
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = p.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("$") and tok.endswith("$") and len(tok) > 2:
            p._p.append(omml(tok[1:-1]))
            continue
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            # 参考文献里的期刊名；避免把乘号当作斜体
            r = p.add_run(tok[1:-1])
            r.italic = True
        else:
            r = p.add_run(tok.replace("\\*", "*"))
        if size:
            r.font.size = Pt(size)
            r.font.name = "Times New Roman"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if color:
            r.font.color.rgb = color


def style_para(p, name, size=None, space_before=None, space_after=None,
               align=None, keep_next=False):
    p.style = name
    pf = p.paragraph_format
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    if align is not None:
        pf.alignment = align
    if keep_next:
        pf.keep_with_next = True
    return p


def unnumber(p):
    """去掉模板标题样式自带的自动编号（用于后置章节标题）。"""
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId"); numId.set(qn("w:val"), "0")
    numPr.append(ilvl); numPr.append(numId)
    pPr.append(numPr)


def add_field(p, instr: str, placeholder: str = "1"):
    r1 = p.add_run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)
    r2 = p.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "
    r2._r.append(it)
    r3 = p.add_run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc)
    r4 = p.add_run(placeholder)
    r5 = p.add_run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end")
    r5._r.append(fc)
    return [r1, r2, r3, r4, r5]


# ---------------------------------------------------------------------------
# 表格
# ---------------------------------------------------------------------------
def set_tbl_borders(table, pad=45):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        '<w:tblBorders %s>'
        '<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        '<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        '<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        '</w:tblBorders>' % nsdecls("w"))
    tblPr.append(borders)
    mar = parse_xml(
        '<w:tblCellMar %s>'
        '<w:top w:w="20" w:type="dxa"/><w:left w:w="%d" w:type="dxa"/>'
        '<w:bottom w:w="20" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/>'
        '</w:tblCellMar>' % (nsdecls("w"), pad, pad))
    tblPr.append(mar)
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def cell_bottom_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        '<w:tcBorders %s><w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/></w:tcBorders>'
        % nsdecls("w"))
    tcPr.append(borders)


def add_table(doc, rows, widths, left_cols=2, pad=45, keep_header_next=False,
              keep_together=False, chain_next=False):
    global _tbl_counter
    _tbl_counter += 1
    ncol = len(rows[0])
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_tbl_borders(t, pad)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        tr = t.rows[-1]._tr
        trPr = tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        for ci, val in enumerate(row):
            cell = cells[ci]
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.style = doc.styles["Els-table-text"]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if ri == 0:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if keep_header_next:
                    p.paragraph_format.keep_with_next = True
                r = p.add_run(val)
                r.bold = True; r.font.size = Pt(8); r.font.name = "Times New Roman"
            else:
                al = WD_ALIGN_PARAGRAPH.LEFT if ci < left_cols else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.alignment = al
                add_inline(p, val, size=8)
            if ri == 0:
                cell_bottom_border(cell)
            if (keep_together and ri < len(rows) - 1) or \
                    (chain_next and ri == len(rows) - 1):
                p.paragraph_format.keep_with_next = True
    # 表后留白
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing = Pt(9)
    if chain_next:
        sp.paragraph_format.keep_with_next = True
    r = sp.add_run(""); r.font.size = Pt(2)
    return t


def parse_table(lines):
    rows = []
    for ln in lines:
        if re.match(r"^\|[\s:\-|]+\|$", ln.strip()):
            continue
        cells = [c.replace("\\|", "|").strip()
                 for c in re.split(r"(?<!\\)\|", ln.strip().strip("|"))]
        rows.append(cells)
    return rows


# ---------------------------------------------------------------------------
# 主构建
# ---------------------------------------------------------------------------
def main():
    src = MD.read_text(encoding="utf-8")
    doc = Document(str(TEMPLATE))

    # 1. 清空正文（保留 sectPr）
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    # 2. 页面：A4 / 2.5 cm / 行号；关闭模板的首页/偶数页页眉
    doc.settings.odd_and_even_pages_header_footer = False
    sec = doc.sections[0]
    sec.different_first_page_header_footer = False
    for ref in list(sec._sectPr.findall(qn("w:headerReference"))):
        if ref.get(qn("w:type")) in ("first", "even"):
            sec._sectPr.remove(ref)
    for el in list(sec._sectPr.findall(qn("w:titlePg"))):
        sec._sectPr.remove(el)
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.header_distance = Cm(1.25)
    sec.footer_distance = Cm(1.25)
    sectPr = sec._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    sectPr.insert_element_before(
        ln, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
        "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
        "w:printerSettings", "w:sectPrChange")

    # 3. 页眉 / 页脚
    hdr = sec.header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    hp.paragraph_format.tab_stops.clear_all()
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run("Beyond the calibration range")
    r.italic = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    r.font.name = "Times New Roman"
    hp.add_run("\t")
    for r in add_field(hp, "PAGE"):
        r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        r.font.name = "Times New Roman"

    ftr = sec.footer
    ftr.is_linked_to_previous = False
    fp = ftr.paragraphs[0]
    for r in list(fp.runs):
        r._r.getparent().remove(r._r)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in add_field(fp, "PAGE"):
        r.font.size = Pt(9); r.font.name = "Times New Roman"

    # 4. 图注字典
    figsec = src.split("## Figure captions")[1].split("## Supplementary material")[0]
    caps = {}
    for m in re.finditer(r"\*\*Fig\. (\d+)\.\*\*\s*(.+?)(?=\n\n\*\*Fig\.|\Z)", figsec, re.S):
        caps[int(m.group(1))] = "**Fig. %s.** %s" % (m.group(1), " ".join(m.group(2).split()))

    # 5. 封面信息
    title = src.splitlines()[0].lstrip("# ").strip()
    p = doc.add_paragraph(style="Els-Title")
    add_inline(p, title)

    p = doc.add_paragraph(style="Els-Author")
    for k, (nm, mark) in enumerate([("Yang Yang", "a"), ("Yuxuan Jiang", "a"),
                                    ("Yalu Cai", "a"), ("Hanping Ke", "b"),
                                    ("Honghao Cai", "a,*")]):
        if k:
            p.add_run(", ")
        p.add_run(nm)
        r = p.add_run(mark)
        r.font.superscript = True

    p = doc.add_paragraph(style="Els-Affiliation")
    r = p.add_run("a")
    r.font.superscript = True
    p.add_run(" School of Science, Jimei University, Xiamen, Fujian Province 361021, China")

    p = doc.add_paragraph(style="Els-Affiliation")
    r = p.add_run("b")
    r.font.superscript = True
    p.add_run(" College of Information Engineering, Ningde Normal University, "
              "Ningde, Fujian Province 352100, China")

    p = doc.add_paragraph(style="Els-Affiliation")
    p.add_run("* Corresponding author: Honghao Cai. E-mail: hhcai@jmu.edu.cn")
    p.add_run().add_break()
    p.add_run("ORCID: https://orcid.org/0000-0002-1870-8061")

    # Abstract (structured: Background / Results / Significance)
    abstract = re.search(r"## Abstract\s*\n\n(.+?)\n\n\*\*Keywords", src, re.S).group(1)
    h = doc.add_paragraph(style="Els-Abstract-head"); h.add_run("Abstract")
    for part in re.split(r"\n\s*\n", abstract.strip()):
        if not part.strip():
            continue
        p = doc.add_paragraph(style="Els-Abstract-text")
        add_inline(p, " ".join(part.split()))

    kw = re.search(r"\*\*Keywords:\*\*\s*(.+)", src).group(1).strip()
    p = doc.add_paragraph(style="Els-keywords")
    r = p.add_run("Keywords: "); r.bold = True
    p.add_run(kw)

    # Highlights
    hi = re.search(r"## Highlights\s*\n\n(.+?)\n\n## Abstract", src, re.S).group(1)
    p = doc.add_paragraph()
    style_para(p, "Normal", size=None, space_before=10, space_after=2, keep_next=True)
    r = p.add_run("Highlights"); r.bold = True; r.font.size = Pt(9)
    for line in [l for l in hi.splitlines() if l.strip().startswith("-")]:
        q = doc.add_paragraph(style="Els-bulletlist")
        q.paragraph_format.space_before = Pt(0)
        q.paragraph_format.space_after = Pt(0)
        add_inline(q, line.strip()[1:].strip(), size=9)

    # 6. 正文
    body_md = src.split("## 1. Introduction")[1].split("## References")[0]
    body_md = "## 1. Introduction" + body_md
    lines = body_md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("@@FIG"):
            n = int(re.search(r"\d+", ln).group())
            path, w = FIG[n]
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_before = Pt(8)
            pic.paragraph_format.space_after = Pt(3)
            pic.paragraph_format.keep_with_next = True
            pic.add_run().add_picture(str(ROOT / path), width=Cm(w))
            cp = doc.add_paragraph(style="Els-caption")
            cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            cp.paragraph_format.space_after = Pt(10)
            add_inline(cp, caps[n])
            i += 1
            continue
        if ln.startswith("## "):
            txt = re.sub(r"^\d+(\.\d+)*\.?\s+", "", ln[3:].strip())
            h = doc.add_paragraph(style="Els-1storder-head")
            if txt in {"CRediT authorship contribution statement", "Data availability",
                       "Declaration of competing interest", "Acknowledgements", "Funding"} \
                    or txt.startswith("Declaration of generative AI"):
                unnumber(h)
            h.add_run(txt)
            i += 1
            continue
        if ln.startswith("### "):
            txt = re.sub(r"^\d+(\.\d+)*\.?\s+", "", ln[4:].strip())
            h = doc.add_paragraph(style="Els-2ndorder-head")
            h.add_run(txt)
            i += 1
            continue
        if ln.startswith("$$"):
            block = ln
            while not block.rstrip().endswith("$$"):
                i += 1
                block += "\n" + lines[i]
            tex = block.strip().strip("$").strip()
            m = re.search(r"\\qquad\s*\((\d+)\)\s*$", tex)
            num = m.group(1) if m else None
            if m:
                tex = tex[:m.start()].rstrip().rstrip(",")
            p = doc.add_paragraph()
            p.paragraph_format.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p.add_run("\t")
            p._p.append(omml(tex))
            if num:
                p.add_run("\t(%s)" % num)
            i += 1
            continue
        if ln.startswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i])
                i += 1
            rows = parse_table(tbl_lines)
            widths = TBL_WIDTHS.get(_tbl_counter + 1, [16.0 / len(rows[0])] * len(rows[0]))
            assert len(widths) == len(rows[0]), (
                "TBL_WIDTHS[%d] has %d entries but table has %d columns"
                % (_tbl_counter + 1, len(widths), len(rows[0])))
            j = i
            while j < len(lines) and not lines[j].strip():
                j += 1
            next_is_table = j < len(lines) and lines[j].strip().startswith("|")
            add_table(doc, rows, widths,
                      keep_together=(len(rows) <= 16 and len(rows[0]) <= 8),
                      chain_next=next_is_table)
            continue
        if ln.startswith("**Table"):
            p = doc.add_paragraph(style="Els-caption")
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            add_inline(p, ln.strip())
            i += 1
            continue
        if ln.startswith("- "):
            p = doc.add_paragraph(style="Els-bulletlist")
            add_inline(p, ln[2:].strip())
            i += 1
            continue
        # 普通段落
        p = doc.add_paragraph(style="Els-body-text")
        add_inline(p, " ".join(ln.split()) if not ln.startswith("**") else ln.strip())
        i += 1

    # 7. References（不编号标题）
    h = doc.add_paragraph(style="Els-1storder-head")
    unnumber(h)
    h.add_run("References")
    refs = src.split("## References")[1].split("## Figure captions")[0]
    for ln in refs.splitlines():
        ln = ln.strip()
        if not ln:
            continue
        p = doc.add_paragraph(style="Els-reference")
        p.paragraph_format.space_after = Pt(2)
        add_inline(p, ln, size=8)

    # 8. 后置章节
    tail = "## Supplementary material" + src.split("## Supplementary material")[1].split("## Figure captions")[0] \
        if "## Supplementary material" in src else ""
    for ln in tail.splitlines():
        ln = ln.strip()
        if not ln:
            continue
        if ln.startswith("## "):
            h = doc.add_paragraph(style="Els-1storder-head")
            unnumber(h)
            h.add_run(ln[3:].strip())
            h.paragraph_format.space_before = Pt(2)
            h.paragraph_format.space_after = Pt(4)
        else:
            p = doc.add_paragraph(style="Els-body-text")
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = Pt(11)
            add_inline(p, ln, size=9)

    doc.core_properties.title = title
    doc.core_properties.author = ""
    doc.core_properties.comments = ""
    doc.save(str(OUT))
    print("saved", OUT)
    print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables),
          "images", len(doc.inline_shapes))
    print("Procedia left?", "Procedia" in "\n".join(p.text for p in doc.paragraphs) or
          "Procedia" in sec.header.paragraphs[0].text)


if __name__ == "__main__":
    main()
