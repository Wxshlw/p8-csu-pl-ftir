"""exp24 —— 外部验证：在两套未参与任何设计与调参的公开数据上原样运行审计与协议。

数据（原始文件不随代码分发，按 DOI 自行下载，放入 P8_EXT_DATA 指向的目录）

E1  红辣椒粉掺 Sudan IV
    Mendeley Data, doi:10.17632/gwgn2th6mg.1 (Brar, Singh, Nanda, 2026, CC BY 4.0)
    文件 FTIR_DATA_Pure_RcP_Sudan(IV)_AdulteratedRcP.xlsx
    按数据说明：ATR-FTIR，Perkin Elmer，700–3700 cm^-1，1.98 cm^-1，透射率（%T），
    7 个品种各 8 条纯样；0.1/0.2/0.4/0.6/1.0 %(w/w) 掺假谱，品种 BM/GM/UT，每水平 4 条。

E2  食用昆虫粉掺玉米/小麦/大豆粉
    Repositório de Dados da UNICAMP, doi:10.25824/redu/zjn5fi (Cruz-Tirado, Barbin, 2026,
    CC BY-NC 4.0)
    文件 "Dataset 1"（xlsx，sheet 'FTIR'）
    按数据说明：FTIR 光谱（未注明是否为 ATR），4000–550 cm^-1，1789 通道，吸光度；
    三种昆虫（buffalo worm/cricket/mealworm）× {0,5,10,20,30,40,50} %(w/w)，每水平 4 条；
    纯样 80 条（28/28/24），掺假样 3 种掺杂物 × 6 水平 × 4 条。

处理约定与正文一致：吸光度；波数升序；E1 由 %T 转 A = -log10(T/100)；
Savitzky–Golay 一阶导数 (window 15, poly 3)；基线阶数 1；λ_s = 0。

输出（写入 out/）
  exp24_ext_predictions.csv  逐数据集×协议×方法的 RMSEP、bias、slope
  exp24_ext_summary.csv      逐数据集×协议：CSU-PL、最优固定比较器、比值
  exp24_ext_audit.json       逐数据集审计量（CNR、CNR_eff、γ、leak、可分辨性、端元余弦）
  exp24b_ext_robust.json     预处理与参考半集敏感性
"""
from __future__ import annotations

import io
import json
import os
import re
import sys
import time
from collections import Counter

import numpy as np
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
OUT = os.path.join(ROOT, "out")
os.makedirs(OUT, exist_ok=True)

from unmix_ftir import baselines, design, diagnostics, forward as F, stats, unmix

SHORT = {"Buffalo worm": "BW", "Cricket": "CR", "Mealworm": "MW",
         "BW": "BW", "CR": "CR", "MW": "MW"}

EXT = os.environ.get("P8_EXT_DATA", os.path.join(ROOT, "ext_data"))
CHILLI = os.path.join(EXT, "FTIR_DATA_Pure_RcP_Sudan(IV)_AdulteratedRcP.xlsx")
INSECT = os.path.join(EXT, "redu_insect1.xlsx")

CFG = unmix.UnmixConfig(baseline_order=1, scatter_sigma=0.0, resolution=np.nan, n_iter=40)
SG = F.SG(window=15, poly=3, deriv=1)
METHODS = ["CSU-PL", "PLS", "Ridge", "CLS-2pt", "MCR-ALS"]
FIXED = METHODS[1:]

PROTOCOLS = {
    "E1": dict(low=([0.2, 0.4, 0.6, 1.0], [0.1]),
               high=([0.0, 0.1, 0.2, 0.4], [0.6, 1.0]),
               interp=([0.0, 0.1, 0.6, 1.0], [0.2, 0.4])),
    "E2": dict(low=([20, 30, 40, 50], [5, 10]),
               high=([0, 5, 10, 20], [40, 50]),
               interp=([0, 5, 10, 40, 50], [20, 30])),
}


# ---------------------------------------------------------------------------
# 读入与规整
# ---------------------------------------------------------------------------
def _prep_insect(path=INSECT):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["FTIR"].iter_rows(values_only=True))
    wvn = np.array([float(v) for v in rows[2][1:-1]])
    X, y, brand, grp = [], [], [], []
    for r in rows[3:]:
        lab = r[0]
        if lab is None:
            continue
        lab = str(lab).strip()
        if lab == "Class" or lab.endswith("dataset"):
            continue
        X.append([float(v) for v in r[1:-1]])
        y.append(float(r[-1]))
        sp, _, ad = lab.replace("  ", " ").partition("+")
        sp = sp.strip()
        ad = ad.strip()
        short = SHORT[sp]
        brand.append(short)
        grp.append(f"{short}|{ad or 'pure'}|{float(r[-1]):g}")
    X = np.asarray(X, float)
    y = np.asarray(y, float)
    if wvn[0] > wvn[-1]:
        wvn = wvn[::-1].copy()
        X = X[:, ::-1].copy()
    return X, y, wvn, np.array(brand, dtype=object), np.array(grp, dtype=object)


def _prep_chilli(path=CHILLI):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Pure_RcP"].iter_rows(values_only=True))
    hp = [h for h in rows[0] if h]
    names = [re.sub(r"\d+", "", str(h).strip()) for h in hp[1:]]
    wvn_p = np.array([float(r[0]) for r in rows[4:]])
    P = np.array([[float(v) for v in r[1:len(hp)]] for r in rows[4:]])
    rows = list(wb["RcP_Adulterted_with_SudanIV"].iter_rows(values_only=True))
    ha = [h for h in rows[0] if h]
    wvn_a = np.array([float(r[0]) for r in rows[2:]])
    A = np.array([[float(v) for v in r[1:len(ha) + 1]] for r in rows[2:]])
    keep, seen = [], set()
    for j, h in enumerate(ha):                      # 分发表含整列重复，按内容去重
        key = (h, A[:, j].tobytes())
        if key in seen:
            continue
        seen.add(key)
        keep.append(j)
    A = A[:, keep]
    ha = [ha[j] for j in keep]
    lo, hi = max(wvn_p.min(), wvn_a.min()), min(wvn_p.max(), wvn_a.max())
    grid = wvn_p[(wvn_p >= lo) & (wvn_p <= hi)]
    P = np.array([np.interp(grid, wvn_p, P[:, j]) for j in range(P.shape[1])]).T
    A = np.array([np.interp(grid, wvn_a, A[:, j]) for j in range(A.shape[1])]).T
    X = [P[:, j] for j in range(P.shape[1])]
    y = [0.0] * P.shape[1]
    brand = list(names)
    grp = [f"{b}|0" for b in names]
    for j, h in enumerate(ha):
        var = h.split("_")[0].replace("Sd", "")
        lvl = float(h.split("_")[1].split("(")[0])
        X.append(A[:, j])
        y.append(lvl)
        brand.append(var)
        grp.append(f"{var}|{lvl:g}|{h[-2]}")
    X = np.asarray(X, float)
    y = np.asarray(y, float)
    X = -np.log10(np.clip(X / 100.0, 1e-4, None))    # %T -> A
    if grid[0] > grid[-1]:
        grid = grid[::-1].copy()
        X = X[:, ::-1].copy()
    return X, y, grid, np.array(brand, dtype=object), np.array(grp, dtype=object)


def load_all():
    """返回 {(数据集, 条件): (X, y, wvn, brand, grp)}"""
    X, y, wvn, brand, grp = _prep_chilli()
    keep = np.isin(brand, ["BM", "GM", "UT"])
    out = {("E1", "Sudan IV"): (X[keep], y[keep], wvn, brand[keep], grp[keep])}
    X, y, wvn, brand, grp = _prep_insect()
    for ad in ("corn", "wheat", "soybean"):
        m = np.isin([g.split("|")[1] for g in grp], [ad, "pure"])
        out[("E2", ad)] = (X[m], y[m], wvn, brand[m], grp[m])
    return out


# ---------------------------------------------------------------------------
# 审计与协议
# ---------------------------------------------------------------------------
def audit(Xv, y, wvn, brand, grp, cfg=CFG):
    mdl = unmix.fit(Xv, y, wvn, cfg, brand=brand)
    curv = diagnostics.curvature_diagnostic(y, Xv)
    em = diagnostics.endmember_audit(mdl, Xv, y)
    rs = diagnostics.label_resolvability(y, Xv, grp)
    zero = Xv[np.isclose(y, 0.0)].mean(0)          # 物理性：各基质端元与零含量均谱的余弦
    zc = zero - zero.mean()
    cos = [float((g - g.mean()) @ zc / (np.linalg.norm(g - g.mean())
                                        * np.linalg.norm(zc) + 1e-18))
           for g in mdl.g[:mdl.n_matrix]]
    gate = "pass" if (em["cnr"] >= 10 and abs(curv["gamma"]) <= 1) else "fail"
    lv = [float(v) for v in sorted(set(np.round(y, 6)))]
    return dict(cnr=float(em["cnr"]), cnr_eff=float(em["cnr_partial"]),
                gamma=float(curv["gamma"]), p_quad=float(curv["p_quad"]),
                leak=float(em["baseline_leak"]), resolv=float(rs["adjacent_snr_med"]),
                cos_endmember=[round(c, 3) for c in cos],
                channels=int(len(wvn)), brands=int(len(np.unique(brand))),
                gate=gate, n=int(len(y)), levels=lv,
                level_counts=[int(np.isclose(y, v).sum()) for v in lv])


def run_protocol(ds, cond, Xv, y, wvn, brand, grp, proto, cfg=CFG, tr_keep=None):
    rows = []
    for p, (tr_lv, te_lv) in proto.items():
        tr = np.isin(np.round(y, 6), np.round(tr_lv, 6))
        te = np.isin(np.round(y, 6), np.round(te_lv, 6))
        if tr_keep is not None:
            tr = tr & tr_keep
            te = te & tr_keep
        if te.sum() == 0 or tr.sum() < 6:
            continue
        mdl = unmix.fit(Xv[tr], y[tr], wvn, cfg, brand=brand[tr])
        pred = mdl.predict_c(Xv[te], brand=brand[te]) * 100.0
        base = dict(dataset=ds, condition=cond, protocol=p, n_tr=int(tr.sum()),
                    n_te=int(te.sum()), kappa=design.kappa(tr_lv),
                    tr_lo=float(min(tr_lv)), tr_hi=float(max(tr_lv)),
                    te_lo=float(min(te_lv)), te_hi=float(max(te_lv)))
        rows.append(dict(method="CSU-PL", **base, **stats.metrics(y[te], pred)))
        for meth in FIXED:
            try:
                q = baselines.fit_predict(meth, Xv[tr], y[tr], grp[tr], Xv[te], wvn, seed=0)
            except Exception as exc:               # 该比较器在本样本量下不成立则跳过
                print(f"    {meth} skipped: {type(exc).__name__}", flush=True)
                continue
            rows.append(dict(method=meth, **base, **stats.metrics(y[te], q)))
    return rows


def summarise(df):
    rows = []
    for (ds, cond, p), g in df.groupby(["dataset", "condition", "protocol"]):
        r = {row["method"]: row for _, row in g.iterrows()}
        if "CSU-PL" not in r:
            continue
        cand = [(r[m]["rmse"], m) for m in FIXED if m in r]
        best, bestm = min(cand)
        rows.append(dict(dataset=ds, condition=cond, protocol=p,
                         cs=float(r["CSU-PL"]["rmse"]), best_rmse=float(best),
                         best_method=bestm,
                         ratio=float(r["CSU-PL"]["rmse"] / best),
                         slope=float(r["CSU-PL"]["slope"]),
                         bias=float(r["CSU-PL"]["bias"])))
    return pd.DataFrame(rows)


def main():
    t0 = time.time()
    data = load_all()
    preds, aut = [], {}
    for (ds, cond), (X, y, wvn, brand, grp) in data.items():
        Xv = SG.transform(X)
        tag = f"{ds}_{cond.replace(' ', '')}"
        a = audit(Xv, y, wvn, brand, grp)
        aut[tag] = a
        print(f"=== {tag}: n={a['n']} CNR={a['cnr']:.1f} (eff {a['cnr_eff']:.1f}) "
              f"gamma={a['gamma']:+.3f} p={a['p_quad']:.1e} leak={a['leak']:.3f} "
              f"-> {a['gate']}", flush=True)
        proto = PROTOCOLS["E1" if ds == "E1" else "E2"]
        for r in run_protocol(ds, cond, Xv, y, wvn, brand, grp, proto):
            preds.append(r)
            print(f"  {r['protocol']:6s} {r['method']:8s} rmse={r['rmse']:7.3f} "
                  f"bias={r['bias']:+7.3f} slope={r['slope']:+6.2f}", flush=True)
    df = pd.DataFrame(preds)
    df.to_csv(os.path.join(OUT, "exp24_ext_predictions.csv"), index=False)
    s = summarise(df)
    s.to_csv(os.path.join(OUT, "exp24_ext_summary.csv"), index=False)
    with io.open(os.path.join(OUT, "exp24_ext_audit.json"), "w", encoding="utf-8") as f:
        json.dump(aut, f, ensure_ascii=False, indent=1)
    out_rng = s[s.protocol.isin(["low", "high"])]
    print("\nmedian ratio, out-of-range, audit-passing systems: "
          f"{out_rng[~out_rng.dataset.eq('E1')].ratio.median():.3f}")
    print("median ratio, out-of-range, failing system: "
          f"{out_rng[out_rng.dataset.eq('E1')].ratio.median():.3f}")
    print(f"({time.time() - t0:.0f} s)")
    _robust(data)


def _robust(data):
    """预处理与参考半集敏感性：结论不随基线阶数、SG 窗宽或参考谱减半而改变。"""
    variants = {
        "base": (CFG, 15, False),
        "baseline3": (unmix.UnmixConfig(baseline_order=3, scatter_sigma=0.0,
                                        resolution=np.nan, n_iter=40), 15, False),
        "sg11": (CFG, 11, False),
        "sg21": (CFG, 21, False),
        "refhalf": (CFG, 15, True),
    }
    rec = {}
    for vname, (cfg, win, half) in variants.items():
        sg = F.SG(window=win, poly=3, deriv=1)
        rows, audits = [], {}
        for (ds, cond), (X, y, wvn, brand, grp) in data.items():
            Xv = sg.transform(X)
            tag = f"{ds}_{cond.replace(' ', '')}"
            keep = None
            if half:
                keep = np.zeros(len(y), bool)
                for lv in np.unique(np.round(y, 6)):
                    idx = np.where(np.isclose(y, lv))[0]
                    keep[idx[0::2]] = True
                a = audit(Xv[keep], y[keep], wvn, brand[keep], grp[keep], cfg)
            else:
                a = audit(Xv, y, wvn, brand, grp, cfg)
            audits[tag] = a
            proto = PROTOCOLS["E1" if ds == "E1" else "E2"]
            rows += run_protocol(ds, cond, Xv, y, wvn, brand, grp, proto,
                                 cfg=cfg, tr_keep=None if keep is None else ~keep)
            print(f"[{vname}] {tag}: {a['gate']} (CNR {a['cnr']:.1f}, γ {a['gamma']:+.3f})",
                  flush=True)
        s = summarise(pd.DataFrame(rows))
        o = s[s.protocol.isin(["low", "high"])]
        rec[vname] = dict(audit=audits, summary=rows,
                          med_pass=float(o[~o.dataset.eq("E1")].ratio.median()),
                          med_fail=float(o[o.dataset.eq("E1")].ratio.median()))
        print(f"[{vname}] out-of-range median ratio: passing {rec[vname]['med_pass']:.3f}, "
              f"failing {rec[vname]['med_fail']:.3f}", flush=True)
    with io.open(os.path.join(OUT, "exp24b_ext_robust.json"), "w", encoding="utf-8") as f:
        json.dump(rec, f, ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main()
