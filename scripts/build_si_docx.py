"""build_si_docx —— 把 SI.md 排成 Elsevier 模板壳的补充材料 .docx。

与正文同样的 A4 / 2.5 cm 版心、三线表、Times New Roman；不启用行号，
页眉为 "Supplementary material" + 页码。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from build_docx import add_field, add_inline, add_table, parse_table, style_para, unnumber  # noqa: E402

SRC = ROOT / "manuscript" / "SI.md"
OUT = ROOT / "manuscript" / "paper8_SI_v1.docx"
TEMPLATE = Path("Elsevier-template.docx")  # local template copy; path sanitised

TBL_WIDTHS = [
    [1.40, 1.30, 5.60, 4.30, 1.10, 1.50],                       # S1
    [1.15, 1.45, 0.75, 1.00, 1.35, 1.30, 1.30, 1.30, 1.30, 1.60, 1.60],  # S2
    [4.60, 4.60, 3.00, 3.00],                                   # S3
    [1.30, 1.20, 1.30, 1.10, 1.10, 1.90, 1.70, 1.60, 1.30, 1.00],  # S4
    [4.60, 1.70, 1.00, 1.30, 1.30, 3.40],                       # S5
    [3.60, 0.85, 1.55, 1.65, 0.95, 1.60, 0.85, 1.55, 1.65, 1.70],  # S6
    [1.20, 1.20, 0.95, 0.95, 3.55, 3.55, 1.40, 1.40, 1.80],     # S7a
    [1.20, 1.20, 0.95, 0.95, 3.55, 3.55, 1.40, 1.40, 1.80],     # S7b
    [1.30, 1.40, 1.35, 1.70, 1.70, 1.35],                       # S8a
    [1.60, 1.40, 2.70, 2.70, 2.90],                             # S8b
    [1.25, 1.45, 3.10, 2.70, 4.00],                             # S8c
    [1.20, 1.05, 1.00, 1.55, 1.55, 0.95, 1.35, 1.20, 1.55, 1.55, 1.30, 1.00],  # S9
    [1.30, 1.50, 1.60, 2.20, 2.20, 2.10, 2.10, 1.50],           # S10
    [1.20, 1.00, 2.90, 1.30, 2.90, 1.10, 1.10, 1.50, 1.10],     # S11
    [1.40, 1.50, 1.40, 2.00, 1.70, 2.00, 2.00, 1.20],           # S12
    [1.60, 2.40, 2.40, 2.20, 2.20, 2.60, 2.60],                 # S13a
    [1.20, 1.20, 1.00, 1.70, 1.80, 1.50, 2.30, 1.50, 2.30, 1.50],  # S13b
    [3.40, 3.15, 3.15, 3.15, 3.15],                             # S13c
    [1.20, 2.60, 2.00, 1.70, 2.10, 1.60, 1.60, 3.20],           # S14a
    [1.30, 1.70, 3.00, 3.20, 2.60, 2.20],                       # S14b
    [4.00, 4.00, 4.00, 4.00],                                   # S14c
    [1.20, 1.20, 0.95, 1.40, 1.50, 0.95, 1.75, 1.55, 1.05, 1.50, 1.70, 1.20],  # S15a
    [1.20, 1.20, 1.30, 1.15, 1.75, 2.00, 1.45, 1.40, 1.70, 1.90],  # S15b
    [1.08, 1.42, 1.05, 0.68, 1.22, 1.52, 1.22, 1.30, 1.22, 0.94, 1.08, 1.09, 1.09, 1.09],  # S16a
    [1.25, 1.50, 1.45, 1.45, 1.65, 4.25, 1.45, 1.35],             # S16b
    [1.60, 1.20, 2.10, 1.90, 2.10, 1.90],                         # S16c
    [1.10, 1.20, 1.60, 1.50, 1.60, 1.60, 2.60, 2.40, 2.40],         # S17a
    [1.20, 1.70, 1.70, 1.40, 1.70, 1.70, 1.70, 1.70, 1.55, 1.40],   # S17b
    [2.60, 3.60, 9.80],                                             # S17c
    [1.30, 4.20, 1.40, 1.20, 1.20, 1.10, 1.40, 1.70, 1.50, 1.00],   # S18a
    [5.20, 3.00, 2.60, 2.60, 2.60],                                 # S18b
    [2.00, 1.50, 1.50, 1.30, 1.30, 1.30, 1.30, 1.30, 1.50, 1.50, 1.50],  # S19a
    [2.60, 3.20, 2.60, 2.50, 2.50, 2.60],                           # S19b
    [3.40, 4.20, 4.20, 4.20],                                       # S19c
    [1.55, 1.50, 0.85, 0.85, 0.85, 0.90, 2.80, 2.20, 1.15, 1.35, 0.85, 1.15],  # S20a
    [2.00, 1.40, 2.10, 2.10, 2.10, 2.10, 2.10, 2.10],               # S20b
    [1.60, 1.70, 2.00, 2.40, 2.20, 2.30, 1.60, 2.20],               # S20c
    [1.60, 1.10, 1.00, 1.70, 1.30, 2.40, 1.70, 1.70, 1.75, 1.75],   # S21
    [3.10, 0.70, 2.20, 1.10, 3.35, 2.75, 1.10],                     # S22
    [1.15, 2.60, 1.15, 1.10, 1.00, 1.30, 1.30, 1.05, 1.20, 0.90, 0.95, 1.35, 0.95],  # S23a
    [1.10, 1.30, 1.70, 1.25, 0.90, 0.90, 1.35, 1.05, 1.30, 1.95, 1.65, 1.10],  # S23b
    [2.20, 1.10, 1.10, 1.10, 1.10, 4.20, 4.20],                     # S23c
    [2.20, 1.50, 3.20, 1.40, 3.40, 1.80, 2.50],                     # S24a
    [1.40, 1.50, 2.70, 2.10, 2.10, 2.30, 2.00, 1.90],               # S24b
    [1.40, 1.50, 2.20, 2.90, 2.70, 2.40, 2.90],                     # S25
    [1.40, 1.50, 2.70, 2.70, 2.60, 2.60, 2.50],                     # S26a
    [1.40, 1.50, 1.40, 2.60, 3.60, 3.30, 2.20],                     # S26b
    [1.60, 1.60, 3.00, 3.10, 3.00, 3.70],                            # S26c
    [3.30, 1.10, 2.30, 2.40, 1.90, 2.30, 2.70],                            # S26d
]
LEFT_COLS = [4, 2, 2, 1, 1, 1, 4, 4, 2, 1, 2, 1, 2, 1, 1,
             1, 3, 1, 2, 2, 1, 2, 2, 3, 2, 1, 1, 1, 3,
             2, 1, 1, 1, 1, 2, 2, 1, 2, 1, 2, 4, 1,
             1, 1, 2, 2, 3, 2, 1]
assert len(TBL_WIDTHS) == len(LEFT_COLS)


def setup(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    doc.settings.odd_and_even_pages_header_footer = False
    sec = doc.sections[0]
    sec.different_first_page_header_footer = False
    for ref in list(sec._sectPr.findall(qn("w:headerReference"))):
        if ref.get(qn("w:type")) in ("first", "even"):
            sec._sectPr.remove(ref)
    for el in list(sec._sectPr.findall(qn("w:titlePg"))):
        sec._sectPr.remove(el)
    for el in list(sec._sectPr.findall(qn("w:lnNumType"))):
        sec._sectPr.remove(el)
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.header_distance = Cm(1.25)
    sec.footer_distance = Cm(1.25)

    hdr = sec.header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]
    for r in list(hp.runs):
        r._r.getparent().remove(r._r)
    hp.paragraph_format.tab_stops.clear_all()
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run("Supplementary material")
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    r.font.name = "Times New Roman"
    hp.add_run("\t")
    for r in add_field(hp, "PAGE"):
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        r.font.name = "Times New Roman"

    ftr = sec.footer
    ftr.is_linked_to_previous = False
    fp = ftr.paragraphs[0]
    for r in list(fp.runs):
        r._r.getparent().remove(r._r)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in add_field(fp, "PAGE"):
        r.font.size = Pt(9)
        r.font.name = "Times New Roman"


def main() -> None:
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document(str(TEMPLATE))
    setup(doc)

    ti = 0
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("# "):
            p = doc.add_paragraph(style="Els-Title")
            r = p.add_run(ln[2:].strip())
            r.font.size = Pt(14)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue
        if ln.startswith("## "):
            p = doc.add_paragraph()
            style_para(p, "Normal", size=11.5, space_before=14, space_after=4, keep_next=True)
            r = p.add_run(ln[3:].strip())
            r.bold = True
            r.font.size = Pt(11.5)
            r.font.name = "Times New Roman"
            if ln.startswith(("## S2.", "## S5.", "## S7.", "## S8.", "## S10.", "## S13.",
                              "## S16.", "## S18.", "## S19.", "## S20.", "## S21.")):
                p.paragraph_format.page_break_before = True
            i += 1
            continue
        if ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = parse_table(block)
            ncol = len(rows[0])
            if ti < len(TBL_WIDTHS) and len(TBL_WIDTHS[ti]) == ncol:
                w, lc = TBL_WIDTHS[ti], LEFT_COLS[ti]
            else:
                w = [16.0 / ncol] * ncol
                lc = 2
                print(f"note: table #{ti + 1} has {ncol} columns, width entry "
                      f"{'missing' if ti >= len(TBL_WIDTHS) else str(len(TBL_WIDTHS[ti]))}; "
                      f"using equal widths")
            add_table(doc, rows, w, left_cols=lc, pad=28 if ncol >= 13 else 45,
                      keep_header_next=True)
            ti += 1
            continue
        if ln.startswith("!!FIG "):
            _, path, wcm = ln.split()
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_before = Pt(8)
            pic.paragraph_format.space_after = Pt(3)
            pic.paragraph_format.keep_with_next = True
            pic.add_run().add_picture(str(ROOT / path), width=Cm(float(wcm)))
            i += 1
            continue
        if ln.startswith("**Table S") or ln.startswith("**Fig. S"):
            p = doc.add_paragraph(style="Els-caption")
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            add_inline(p, ln.strip())
            i += 1
            continue
        if ln.startswith("**Beyond the calibration range"):
            p = doc.add_paragraph()
            style_para(p, "Normal", size=10, space_before=0, space_after=10,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
            add_inline(p, ln.strip(), size=10)
            i += 1
            continue
        if ln.startswith("Note:"):
            p = doc.add_paragraph()
            style_para(p, "Normal", size=7.5, space_before=3, space_after=10)
            add_inline(p, ln.strip(), size=7.5)
            i += 1
            continue
        if re.match(r"^\([abc]\)\s", ln):
            p = doc.add_paragraph()
            style_para(p, "Normal", size=9, space_before=6, space_after=2,
                       keep_next=True)
            r = p.add_run(ln.strip())
            r.bold = True
            r.font.size = Pt(9)
            r.font.name = "Times New Roman"
            i += 1
            continue
        # 普通段落
        p = doc.add_paragraph()
        style_para(p, "Els-body-text", size=9.5, space_before=0, space_after=4)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, ln.strip(), size=9.5)
        i += 1

    doc.save(str(OUT))
    print(f"saved {OUT}")


if __name__ == "__main__":
    main()
